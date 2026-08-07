from openpyxl import load_workbook
import os
from utils.logger import logger


def parse_excel(ruta_archivo: str, hoja: str | None = None) -> tuple[bool, str, list[dict]]:
    if not os.path.exists(ruta_archivo):
        return False, "El archivo no existe", []

    extension = os.path.splitext(ruta_archivo)[1].lower()
    if extension not in (".xlsx", ".xls"):
        return False, "El archivo no es un Excel válido", []

    try:
        wb = load_workbook(ruta_archivo, read_only=True, data_only=True)

        if hoja and hoja in wb.sheetnames:
            ws = wb[hoja]
        else:
            ws = wb.active

        if ws is None:
            return False, "No se encontró ninguna hoja en el archivo", []

        filas = []
        encabezados = None

        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row is None or all(c is None for c in row):
                continue

            valores = [str(c).strip() if c is not None else "" for c in row]

            if encabezados is None:
                encabezados = valores
                continue

            if all(v == "" for v in valores):
                continue

            fila = {}
            for i, h in enumerate(encabezados):
                if i < len(valores):
                    fila[h] = valores[i]
                else:
                    fila[h] = ""
            filas.append(fila)

        wb.close()

        if not encabezados:
            return False, "El archivo está vacío o no tiene encabezados", []

        if not filas:
            return False, "El archivo no tiene datos", []

        logger.info(f"Excel parseado: {len(filas)} filas, {len(encabezados)} columnas")
        return True, f"{len(filas)} registros encontrados", filas

    except Exception as e:
        logger.error(f"Error al parsear Excel: {e}")
        return False, f"Error al leer el archivo: {str(e)}", []


def obtener_hojas_excel(ruta_archivo: str) -> tuple[bool, str, list[str]]:
    if not os.path.exists(ruta_archivo):
        return False, "El archivo no existe", []

    try:
        wb = load_workbook(ruta_archivo, read_only=True)
        hojas = wb.sheetnames
        wb.close()
        return True, "OK", hojas

    except Exception as e:
        logger.error(f"Error al leer hojas Excel: {e}")
        return False, f"Error: {str(e)}", []


def obtener_columnas_excel(ruta_archivo: str, hoja: str | None = None) -> tuple[bool, str, list[str]]:
    if not os.path.exists(ruta_archivo):
        return False, "El archivo no existe", []

    try:
        wb = load_workbook(ruta_archivo, read_only=True, data_only=True)

        if hoja and hoja in wb.sheetnames:
            ws = wb[hoja]
        else:
            ws = wb.active

        if ws is None:
            wb.close()
            return False, "No se encontró ninguna hoja", []

        for row in ws.iter_rows(values_only=True):
            if row and any(c is not None for c in row):
                columnas = [str(c).strip() if c is not None else "" for c in row]
                wb.close()
                return True, "OK", columnas

        wb.close()
        return False, "No se encontraron encabezados", []

    except Exception as e:
        logger.error(f"Error al leer columnas Excel: {e}")
        return False, f"Error: {str(e)}", []
