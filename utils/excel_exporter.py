from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os


def exportar_a_excel(datos: list[dict], columnas: list[tuple[str, str]],
                     titulo: str, ruta_archivo: str) -> tuple[bool, str]:
    if not datos:
        return False, "No hay datos para exportar"

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    title_cell = ws.cell(row=1, column=1, value=titulo)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    ws.cell(row=2, column=1, value=f"Fecha: {__import__('datetime').date.today().strftime('%d/%m/%Y')}")

    for col_idx, (key, header_text) in enumerate(columnas, 1):
        cell = ws.cell(row=4, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, fila in enumerate(datos, 5):
        for col_idx, (key, _) in enumerate(columnas, 1):
            valor = fila.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left")

    for col_idx, _ in enumerate(columnas, 1):
        max_length = max(
            len(str(ws.cell(row=row, column=col_idx).value or ""))
            for row in range(4, len(datos) + 6)
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = max(max_length + 2, 12)

    try:
        os.makedirs(os.path.dirname(ruta_archivo), exist_ok=True)
        wb.save(ruta_archivo)
        return True, f"Reporte exportado: {ruta_archivo}"
    except Exception as e:
        return False, f"Error al exportar: {str(e)}"
